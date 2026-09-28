from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from agent.llm import MockLLMClient
from agent.prompts import TOOL_SCHEMAS
from agent.schemas import ToolCall
from app import create_app
from models import (
    AccountPayable,
    AccountReceivable,
    Customer,
    InventoryTransaction,
    Product,
    PurchaseOrder,
    PurchaseOrderItem,
    SalesOrder,
    SalesOrderItem,
    Supplier,
    db,
)


@pytest.fixture()
def app(tmp_path: Path):
    app = create_app(
        {
            "TESTING": True,
            "SECRET_KEY": "test-secret",
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'agent_routes.db'}",
            "AGENT_LLM_CLIENT": MockLLMClient([]),
        }
    )

    with app.app_context():
        db.create_all()

    yield app

    with app.app_context():
        db.session.remove()
        db.drop_all()


@pytest.fixture()
def master_data(app):
    with app.app_context():
        supplier = Supplier(name="南京键盘供应商", phone="")
        customer = Customer(name="xx科技", phone="")
        keyboard = Product(
            name="机械键盘",
            sku="KB001",
            purchase_price=80,
            sale_price=120,
            stock=80,
        )
        db.session.add_all([supplier, customer, keyboard])
        db.session.commit()
        return {
            "supplier_id": supplier.id,
            "customer_id": customer.id,
            "product_id": keyboard.id,
        }


def set_llm(app, *responses):
    client = MockLLMClient(list(responses))
    app.config["AGENT_LLM_CLIENT"] = client
    return client


def purchase_call(product_name="机械键盘"):
    return ToolCall(
        name="prepare_purchase_order",
        arguments={
            "supplier_name": "南京键盘供应商",
            "items": [
                {
                    "product_name": product_name,
                    "sku": None,
                    "quantity": 2,
                    "unit_price": 78,
                }
            ],
        },
    )


def sales_call():
    return ToolCall(
        name="prepare_sales_order",
        arguments={
            "customer_name": "xx科技",
            "items": [
                {
                    "product_name": "机械键盘",
                    "sku": None,
                    "quantity": 100,
                    "unit_price": 120,
                }
            ],
        },
    )


def unpaid_receivables_call(customer_name=None):
    return ToolCall(
        name="get_unpaid_receivables",
        arguments={"customer_name": customer_name},
    )


def seed_unpaid_receivable(app, customer_id, order_no="SO-UNPAID-ROUTE"):
    with app.app_context():
        order = SalesOrder(
            order_no=order_no,
            customer_id=customer_id,
            status="completed",
            total_amount=240,
        )
        db.session.add(order)
        db.session.flush()
        receivable = AccountReceivable(
            sales_order_id=order.id,
            customer_id=customer_id,
            amount=240,
            status="unpaid",
        )
        db.session.add(receivable)
        db.session.commit()
        return receivable.id


@pytest.fixture()
def replenishment_data(app, master_data):
    with app.app_context():
        product = db.session.get(Product, master_data["product_id"])
        product.stock = 20
        sales_order = SalesOrder(
            order_no="SO-REPLENISHMENT-ROUTE",
            customer_id=master_data["customer_id"],
            status="completed",
            created_at=datetime.now(timezone.utc) - timedelta(days=3),
        )
        sales_order.items.append(
            SalesOrderItem(product_id=product.id, quantity=35, unit_price=120)
        )
        db.session.add(sales_order)
        db.session.commit()
        return master_data


def replenishment_call(product_name="机械键盘", supplier_name="南京键盘供应商", quantity=None):
    return ToolCall(
        name="prepare_replenishment_purchase",
        arguments={
            "product_name": product_name,
            "supplier_name": supplier_name,
            "quantity": quantity,
        },
    )


def test_assistant_page_returns_200(app):
    response = app.test_client().get("/assistant")

    assert response.status_code == 200
    assert "AI ERP Assistant" in response.get_data(as_text=True)
    assert "查一下机械键盘库存" in response.get_data(as_text=True)


def test_assistant_exports_all_data_with_shared_filter_arguments(app):
    export_schema = next(
        tool["function"] for tool in TOOL_SCHEMAS
        if tool["function"]["name"] == "export_dataset"
    )
    properties = export_schema["parameters"]["properties"]
    assert "all" in properties["dataset"]["enum"]
    assert export_schema["parameters"]["required"] == [
        "dataset", "file_format", "start_date", "end_date", "limit"
    ]
    assert properties["limit"]["type"] == ["integer", "null"]

    set_llm(
        app,
        ToolCall(
            name="export_dataset",
            arguments={
                "dataset": "all",
                "file_format": "xlsx",
                "start_date": None,
                "end_date": None,
                "limit": 2,
            },
        ),
    )
    client = app.test_client()
    response = client.post(
        "/assistant/message", json={"message": "导出全部数据为 Excel"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    assert response.json["data"]["filename"].startswith("all_data_")
    download = client.get(response.json["data"]["download_url"])
    workbook = load_workbook(BytesIO(download.data), read_only=True)
    assert download.status_code == 200
    assert len(workbook.sheetnames) == 9


def test_assistant_message_returns_verified_inventory_result(app, master_data):
    set_llm(
        app,
        ToolCall(
            name="get_inventory",
            arguments={"product_name": "机械键盘", "sku": None},
        ),
    )

    response = app.test_client().post(
        "/assistant/message", json={"message": "查一下机械键盘库存"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    assert "80" in response.json["content"]


@pytest.mark.parametrize(
    "message",
    [
        "现在还有哪些客户没付款",
        "现在还有哪些客户没有付款",
        "帮我查一下哪些客户没付款",
    ],
)
def test_assistant_routes_unpaid_customer_question_to_read_tool(
    app, master_data, message
):
    receivable_id = seed_unpaid_receivable(app, master_data["customer_id"])
    llm = set_llm(app, unpaid_receivables_call())

    response = app.test_client().post(
        "/assistant/message", json={"message": message}
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    assert response.json["data"]["items"][0]["receivable_id"] == receivable_id
    assert response.json["data"]["items"][0]["customer"] == "xx科技"
    assert response.json["data"]["items"][0]["sales_order_no"] == "SO-UNPAID-ROUTE"
    assert response.json["data"]["items"][0]["amount"] == "240.00"
    assert "xx科技" in response.json["content"]
    assert len(llm.calls) == 1
    with app.app_context():
        assert AccountReceivable.query.one().status == "unpaid"


def test_assistant_routes_named_customer_debt_question_to_read_tool(app, master_data):
    seed_unpaid_receivable(app, master_data["customer_id"])
    llm = set_llm(app, unpaid_receivables_call(customer_name="xx科技"))

    response = app.test_client().post(
        "/assistant/message", json={"message": "xx科技还欠多少钱"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    assert response.json["data"]["total_amount"] == "240.00"
    assert response.json["data"]["items"][0]["customer"] == "xx科技"
    assert len(llm.calls) == 1


def test_purchase_message_returns_preview_without_creating_order(app, master_data):
    set_llm(app, purchase_call())

    response = app.test_client().post(
        "/assistant/message", json={"message": "向南京键盘供应商采购 2 个机械键盘"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "confirmation"
    assert response.json["preview"]["total_amount"] == "156.00"
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


def test_purchase_confirmation_creates_only_after_explicit_confirm(app, master_data):
    set_llm(app, purchase_call())
    client = app.test_client()

    preview_response = client.post(
        "/assistant/message", json={"message": "生成采购预览"}
    )
    token = preview_response.json["confirmation_token"]
    confirm_response = client.post(
        "/assistant/confirm",
        json={"confirmation_token": token, "action": "confirm"},
    )

    assert confirm_response.status_code == 200
    assert confirm_response.json["type"] == "message"
    with app.app_context():
        assert PurchaseOrder.query.count() == 1
        assert PurchaseOrder.query.one().status == "draft"
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


def test_cancelled_purchase_preview_does_not_create_order(app, master_data):
    set_llm(app, purchase_call())
    client = app.test_client()
    preview_response = client.post(
        "/assistant/message", json={"message": "生成采购预览"}
    )

    response = client.post(
        "/assistant/confirm",
        json={
            "confirmation_token": preview_response.json["confirmation_token"],
            "action": "cancel",
        },
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    with app.app_context():
        assert PurchaseOrder.query.count() == 0


def test_invalid_confirmation_token_does_not_create_order(app, master_data):
    response = app.test_client().post(
        "/assistant/confirm",
        json={"confirmation_token": "forged", "action": "confirm"},
    )

    assert response.status_code == 400
    assert response.json["type"] == "error"
    with app.app_context():
        assert PurchaseOrder.query.count() == 0


def test_sales_confirmation_allows_insufficient_stock_and_creates_no_side_effects(
    app, master_data
):
    set_llm(app, sales_call())
    client = app.test_client()
    preview_response = client.post(
        "/assistant/message", json={"message": "给xx科技销售 100 个机械键盘"}
    )

    confirm_response = client.post(
        "/assistant/confirm",
        json={
            "confirmation_token": preview_response.json["confirmation_token"],
            "action": "confirm",
        },
    )

    assert confirm_response.status_code == 200
    with app.app_context():
        assert SalesOrder.query.count() == 1
        assert SalesOrder.query.one().status == "draft"
        assert db.session.get(Product, master_data["product_id"]).stock == 80
        assert InventoryTransaction.query.count() == 0
        assert AccountReceivable.query.count() == 0


def test_replenishment_analysis_returns_real_metrics_and_exposes_tool_schemas(
    app, replenishment_data
):
    llm = set_llm(
        app,
        ToolCall(name="analyze_low_stock", arguments={"product_name": "机械键盘"}),
    )

    response = app.test_client().post(
        "/assistant/message", json={"message": "分析机械键盘是否需要补货"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "message"
    assert response.json["data"]["analysis_type"] == "replenishment"
    assert response.json["data"]["items"][0]["current_stock"] == 20
    assert response.json["data"]["items"][0]["sales_7d"] == 35
    assert response.json["data"]["items"][0]["sales_30d"] == 35
    assert response.json["data"]["items"][0]["recommended_purchase_qty"] == 50

    schemas = {tool["function"]["name"]: tool["function"] for tool in llm.calls[0]["tools"]}
    assert {"analyze_low_stock", "prepare_replenishment_purchase"} <= schemas.keys()
    assert schemas["analyze_low_stock"]["strict"] is True
    replenishment_schema = schemas["prepare_replenishment_purchase"]["parameters"]
    assert replenishment_schema["required"] == [
        "product_name",
        "supplier_name",
        "quantity",
    ]
    assert replenishment_schema["additionalProperties"] is False
    assert replenishment_schema["properties"] == {
        "product_name": {"type": ["string", "null"]},
        "supplier_name": {"type": ["string", "null"]},
        "quantity": {"type": ["integer", "null"], "minimum": 1},
    }


def test_replenishment_preview_and_confirmation_are_safe(app, replenishment_data):
    set_llm(app, replenishment_call())
    client = app.test_client()

    preview_response = client.post(
        "/assistant/message", json={"message": "按建议数量生成机械键盘补货预览"}
    )

    assert preview_response.status_code == 200
    assert preview_response.json["type"] == "confirmation"
    assert preview_response.json["preview"]["items"][0]["quantity"] == 50
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0

    confirm_response = client.post(
        "/assistant/confirm",
        json={
            "confirmation_token": preview_response.json["confirmation_token"],
            "action": "confirm",
        },
    )

    assert confirm_response.status_code == 200
    with app.app_context():
        assert PurchaseOrder.query.count() == 1
        assert PurchaseOrder.query.one().status == "draft"
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


def test_replenishment_preview_defaults_the_only_supplier_without_writing(
    app, replenishment_data
):
    set_llm(
        app,
        replenishment_call(supplier_name=None),
    )

    response = app.test_client().post(
        "/assistant/message", json={"message": "生成机械键盘补货预览"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "confirmation"
    assert response.json["preview"]["supplier_name"] == "南京键盘供应商"
    assert response.json["preview"]["supplier_source"] == "only_supplier"
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


def test_replenishment_context_identifies_product_without_reusing_analysis_facts(
    app, replenishment_data
):
    llm = set_llm(
        app,
        ToolCall(name="analyze_low_stock", arguments={"product_name": "机械键盘"}),
        replenishment_call(product_name=None, supplier_name=None, quantity=40),
    )
    client = app.test_client()

    analysis_response = client.post(
        "/assistant/message", json={"message": "分析机械键盘是否需要补货"}
    )
    preview_response = client.post(
        "/assistant/message", json={"message": "那就补 40 个"}
    )

    assert analysis_response.status_code == 200
    assert preview_response.status_code == 200
    assert preview_response.json["type"] == "confirmation"
    assert preview_response.json["preview"]["items"][0]["product_name"] == "机械键盘"
    assert preview_response.json["preview"]["items"][0]["quantity"] == 40
    assert "机械键盘" in llm.calls[1]["message"]
    assert "仅用于识别商品" in llm.calls[1]["message"]
    assert "不得信任" in llm.calls[1]["message"]
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


@pytest.mark.parametrize("supplier_name", [None, "供应商", "不存在供应商"])
def test_replenishment_supplier_clarification_preserves_active_product_context(
    app, replenishment_data, supplier_name
):
    with app.app_context():
        product = Product(
            name="无线鼠标", sku="MOUSE001", purchase_price=10, sale_price=20, stock=3
        )
        supplier = Supplier(name="北京鼠标供应商", phone="")
        db.session.add_all([product, supplier])
        db.session.flush()
        order = SalesOrder(
            order_no="SO-ACTIVE-PRODUCT",
            customer_id=replenishment_data["customer_id"],
            status="completed",
            created_at=datetime.now(timezone.utc) - timedelta(days=1),
        )
        order.items.append(
            SalesOrderItem(product_id=product.id, quantity=7, unit_price=20)
        )
        db.session.add(order)
        db.session.commit()
        product_id, supplier_id = product.id, supplier.id

    llm = set_llm(
        app,
        ToolCall(name="analyze_low_stock", arguments={"product_name": "机械键盘"}),
        replenishment_call(product_name="无线鼠标", supplier_name=supplier_name),
        replenishment_call(product_name=None, supplier_name="北京鼠标供应商"),
    )
    client = app.test_client()
    analysis = client.post(
        "/assistant/message", json={"message": "分析机械键盘是否需要补货"}
    )
    clarification = client.post(
        "/assistant/message", json={"message": "生成无线鼠标补货预览"}
    )
    assert analysis.status_code == clarification.status_code == 200
    assert clarification.json["type"] == "clarification"
    with client.session_transaction() as session:
        active_context = session["assistant_last_replenishment_context"]

    # Facts change while the user is choosing a supplier; only identity may persist.
    with app.app_context():
        product = db.session.get(Product, product_id)
        product.stock = 4
        product.purchase_price = 12
        order = SalesOrder.query.filter_by(order_no="SO-ACTIVE-PRODUCT").one()
        order.items[0].quantity = 14
        purchase = PurchaseOrder(
            order_no="PO-ACTIVE-PRODUCT",
            supplier_id=supplier_id,
            status="pending_receipt",
        )
        purchase.items.append(
            PurchaseOrderItem(product_id=product_id, quantity=5, unit_price=9)
        )
        db.session.add(purchase)
        db.session.commit()

    response = client.post(
        "/assistant/message", json={"message": "选择北京鼠标供应商"}
    )
    assert response.status_code == 200
    assert response.json["type"] == "confirmation"
    preview = response.json["preview"]
    assert preview["items"][0] == {
        "product_name": "无线鼠标",
        "sku": "MOUSE001",
        "quantity": 19,
        "unit_price": "9.00",
        "subtotal": "171.00",
    }
    assert preview["supplier_name"] == "北京鼠标供应商"
    assert preview["recommendation"]["current_stock"] == 4
    assert preview["recommendation"]["sales_7d"] == 14
    assert preview["recommendation"]["pending_purchase_qty"] == 5
    assert active_context == [{"product_name": "无线鼠标", "sku": "MOUSE001"}]
    assert clarification.json["data"] == {
        "analysis_type": "replenishment",
        "items": active_context,
    }
    assert "无线鼠标" in llm.calls[2]["message"]
    assert "机械键盘" not in llm.calls[2]["message"]
    with app.app_context():
        assert PurchaseOrder.query.count() == 1
        assert PurchaseOrder.query.one().status == "pending_receipt"
        assert InventoryTransaction.query.count() == 0
        assert AccountPayable.query.count() == 0


@pytest.mark.parametrize("tool_name", ["create_purchase_order", "delete_inventory"])
def test_assistant_rejects_unapproved_write_tools_without_mutation(
    app, master_data, tool_name
):
    set_llm(app, ToolCall(name=tool_name, arguments={}))

    response = app.test_client().post(
        "/assistant/message", json={"message": "执行写操作"}
    )

    assert response.status_code == 200
    assert response.json["type"] == "error"
    assert "不支持工具" in response.json["message"]
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert SalesOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountReceivable.query.count() == 0
        assert AccountPayable.query.count() == 0


@pytest.mark.parametrize(
    "message",
    [
        "把 PO20260921001 入库",
        "把 SO20260921001 发货",
        "确认收款",
        "确认付款",
        "执行付款",
        "支付给xx科技",
        "把应收账款收一下",
        "帮我把xx科技的应收收一下",
    ],
)
def test_assistant_refuses_high_risk_operations_without_mutation(app, message):
    response = app.test_client().post("/assistant/message", json={"message": message})

    assert response.status_code == 200
    assert response.json["type"] == "error"
    assert "暂不支持" in response.json["message"]
    assert app.config["AGENT_LLM_CLIENT"].calls == []
    with app.app_context():
        assert PurchaseOrder.query.count() == 0
        assert SalesOrder.query.count() == 0
        assert InventoryTransaction.query.count() == 0
        assert AccountReceivable.query.count() == 0
        assert AccountPayable.query.count() == 0
