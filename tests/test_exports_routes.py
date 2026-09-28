from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app import create_app
from models import Product, db
from services.exports import XLSX_MIME_TYPE


@pytest.fixture()
def app(tmp_path: Path):
    app = create_app(
        {
            "TESTING": True,
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'exports_routes.db'}",
        }
    )
    with app.app_context():
        db.create_all()
    yield app
    with app.app_context():
        db.session.remove()
        db.drop_all()


def test_unified_export_page_lists_datasets_and_snapshot_rule(app):
    response = app.test_client().get("/exports/")
    body = response.get_data(as_text=True)

    assert response.status_code == 200
    assert "全部数据" in body
    assert 'value="inventory_transactions"' in body
    assert 'value="receivables"' in body
    assert "日期范围只筛选有创建时间的数据" in body
    assert "商品、客户、供应商和当前库存" in body


def test_unified_export_post_returns_a_combined_workbook(app):
    response = app.test_client().post(
        "/exports/",
        data={
            "dataset": "all",
            "file_format": "xlsx",
            "start_date": "2026-09-28",
            "end_date": "",
            "limit": "25",
        },
    )

    assert response.status_code == 200
    assert response.mimetype == XLSX_MIME_TYPE
    assert "attachment" in response.headers["Content-Disposition"]
    workbook = load_workbook(BytesIO(response.data), read_only=True)
    assert len(workbook.sheetnames) == 9


def test_existing_direct_export_route_still_exports_the_full_dataset(app):
    with app.app_context():
        db.session.add(
            Product(name="快捷导出商品", sku="FAST-1", purchase_price=1, sale_price=2)
        )
        db.session.commit()

    response = app.test_client().get("/exports/products/xlsx")

    workbook = load_workbook(BytesIO(response.data), read_only=True)
    assert response.status_code == 200
    assert workbook.sheetnames == ["商品列表"]
    assert workbook.active["A3"].value == "快捷导出商品"


@pytest.mark.parametrize(
    ("overrides", "message"),
    [
        ({"start_date": "2026-09-29", "end_date": "2026-09-28"}, "开始日期不能晚于结束日期"),
        ({"limit": "10001"}, "条数必须是 1 至 10,000 的整数"),
        ({"dataset": "user_table"}, "不支持导出该数据集"),
        ({"file_format": "csv"}, "导出格式仅支持 Excel 或 PDF"),
    ],
)
def test_invalid_export_post_renders_validation_error(app, overrides, message):
    data = {
        "dataset": "all",
        "file_format": "xlsx",
        "start_date": "",
        "end_date": "",
        "limit": "",
    }
    data.update(overrides)

    response = app.test_client().post("/exports/", data=data)

    body = response.get_data(as_text=True)
    assert response.status_code == 400
    assert message in body
    assert 'method="post"' in body
    assert f'name="dataset"' in body

