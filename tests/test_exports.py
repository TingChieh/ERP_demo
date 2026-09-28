from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app import create_app
from models import Customer, Product, SalesOrder, db
from services import exports


@pytest.fixture()
def app(tmp_path: Path):
    app = create_app(
        {
            "TESTING": True,
            "SQLALCHEMY_DATABASE_URI": f"sqlite:///{tmp_path / 'exports.db'}",
        }
    )
    with app.app_context():
        db.create_all()
    yield app
    with app.app_context():
        db.session.remove()
        db.drop_all()


def test_parse_export_filters_converts_inclusive_shanghai_days_to_utc():
    filters = exports.parse_export_filters("2026-09-28", "2026-09-29", "5")

    assert filters.start_at == datetime(2026, 9, 27, 16, 0)
    assert filters.end_before == datetime(2026, 9, 29, 16, 0)
    assert filters.limit == 5


def test_parse_export_filters_keeps_a_single_date_boundary():
    filters = exports.parse_export_filters(None, "2026-09-28", None)

    assert filters.start_at is None
    assert filters.end_before == datetime(2026, 9, 28, 16, 0)
    assert filters.limit is None


@pytest.mark.parametrize(
    ("start_date", "end_date", "limit"),
    [
        ("2026-09-29", "2026-09-28", None),
        ("28-09-2026", None, None),
        (None, None, 0),
        (None, None, 10_001),
        (None, None, True),
        (None, None, "1.5"),
    ],
)
def test_parse_export_filters_rejects_invalid_values(start_date, end_date, limit):
    with pytest.raises(exports.ExportRequestError):
        exports.parse_export_filters(start_date, end_date, limit)


def test_sales_orders_use_inclusive_date_bounds_and_latest_first_limit(app):
    with app.app_context():
        customer = Customer(name="边界客户", phone="")
        db.session.add(customer)
        db.session.flush()
        before = SalesOrder(
            order_no="SO-BEFORE",
            customer_id=customer.id,
            status="draft",
            total_amount=0,
            created_at=datetime(2026, 9, 27, 15, 59, 59, 999999, tzinfo=timezone.utc),
        )
        at_start = SalesOrder(
            order_no="SO-START",
            customer_id=customer.id,
            status="draft",
            total_amount=0,
            created_at=datetime(2026, 9, 27, 16, 0, tzinfo=timezone.utc),
        )
        at_end = SalesOrder(
            order_no="SO-END",
            customer_id=customer.id,
            status="draft",
            total_amount=0,
            created_at=datetime(2026, 9, 28, 15, 59, 59, 999999, tzinfo=timezone.utc),
        )
        after = SalesOrder(
            order_no="SO-AFTER",
            customer_id=customer.id,
            status="draft",
            total_amount=0,
            created_at=datetime(2026, 9, 28, 16, 0, tzinfo=timezone.utc),
        )
        db.session.add_all([before, at_start, at_end, after])
        db.session.commit()

        filters = exports.parse_export_filters("2026-09-28", "2026-09-28", 1)
        _, rows = exports.get_export_rows("sales_orders", filters)

        assert [row[0] for row in rows] == ["SO-END"]


def test_snapshot_dataset_ignores_dates_but_applies_limit(app):
    with app.app_context():
        db.session.add_all(
            [
                Product(name="旧商品", sku="P-OLD", purchase_price=1, sale_price=1),
                Product(name="新商品", sku="P-NEW", purchase_price=1, sale_price=1),
            ]
        )
        db.session.commit()
        filters = exports.parse_export_filters("2000-01-01", "2000-01-01", 1)

        specification, rows = exports.get_export_rows("products", filters)

        assert specification.supports_date_filter is False
        assert [row[1] for row in rows] == ["P-NEW"]


def test_all_xlsx_has_one_sheet_for_each_dataset_in_fixed_order(app):
    with app.app_context():
        export_file = exports.generate_export("all", "xlsx")
        workbook = load_workbook(BytesIO(export_file.content), read_only=True)

        assert workbook.sheetnames == [
            dataset.title for dataset in exports.DATASETS.values()
        ]
        assert export_file.mime_type == exports.XLSX_MIME_TYPE
        assert export_file.filename.startswith("all_data_")


def test_empty_xlsx_sheet_includes_a_clear_empty_result_message(app):
    with app.app_context():
        export_file = exports.generate_export("products", "xlsx")
        workbook = load_workbook(BytesIO(export_file.content), read_only=True)

        assert workbook.active["A3"].value == "没有符合条件的记录"


def test_all_pdf_is_one_report_with_sections_and_snapshot_notes(app):
    with app.app_context():
        export_file = exports.generate_export("all", "pdf")
        reader = PdfReader(BytesIO(export_file.content))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)

        assert export_file.content.startswith(b"%PDF-")
        assert export_file.mime_type == exports.PDF_MIME_TYPE
        assert export_file.filename.startswith("all_data_")
        assert all(dataset.title in text for dataset in exports.DATASETS.values())
        assert "当前快照，日期范围不适用" in text
        assert "商品名称" in text
        assert "没有符合条件的记录" in text


def test_single_snapshot_pdf_discloses_that_dates_do_not_apply(app):
    with app.app_context():
        export_file = exports.generate_export(
            "products", "pdf", start_date="2026-09-28", end_date="2026-09-28"
        )
        text = "\n".join(
            page.extract_text() or "" for page in PdfReader(BytesIO(export_file.content)).pages
        )

        assert "当前快照，日期范围不适用" in text
