from dataclasses import dataclass
from datetime import date, datetime, time, timedelta, timezone
from html import escape
from io import BytesIO
from typing import Callable
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Table,
    TableStyle,
)
from sqlalchemy.orm import joinedload, selectinload

from models import (
    AccountPayable,
    AccountReceivable,
    Customer,
    InventoryTransaction,
    Product,
    PurchaseOrder,
    PurchaseOrderItem,
    SalesOrder,
    SalesOrderItem,
    Supplier,
)


XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MIME_TYPE = "application/pdf"
MONEY_FORMAT = '¥#,##0.00;[Red]-¥#,##0.00'
ALL_DATASET_ID = "all"
STATUS_LABELS = {
    "draft": "草稿",
    "pending_receipt": "待入库",
    "pending_shipment": "待出库",
    "completed": "已完成",
    "unpaid": "未结清",
    "paid": "已结清",
}
MOVEMENT_LABELS = {"inbound": "入库", "outbound": "出库"}
MAX_EXPORT_ROWS = 10_000
EXPORT_TIME_ZONE = ZoneInfo("Asia/Shanghai")


class ExportRequestError(ValueError):
    """Raised when the export dataset or file format is not supported."""


@dataclass(frozen=True)
class ExportFilters:
    start_at: datetime | None = None
    end_before: datetime | None = None
    limit: int | None = None


@dataclass(frozen=True)
class ExportColumn:
    label: str
    kind: str = "text"


@dataclass(frozen=True)
class ExportDataset:
    dataset_id: str
    title: str
    columns: tuple[ExportColumn, ...]
    query_rows: Callable[[ExportFilters], list[tuple[object, ...]]]
    supports_date_filter: bool = False


@dataclass(frozen=True)
class ExportFile:
    content: bytes
    filename: str
    mime_type: str


def _parse_export_date(label, value):
    if value is None or value == "":
        return None
    if not isinstance(value, str):
        raise ExportRequestError(f"{label}格式应为 YYYY-MM-DD。")
    raw = value.strip()
    try:
        parsed = date.fromisoformat(raw)
    except ValueError:
        raise ExportRequestError(f"{label}格式应为 YYYY-MM-DD。") from None
    if parsed.isoformat() != raw:
        raise ExportRequestError(f"{label}格式应为 YYYY-MM-DD。")
    return parsed


def _local_midnight_to_utc_naive(day):
    return datetime.combine(day, time.min, tzinfo=EXPORT_TIME_ZONE).astimezone(
        timezone.utc
    ).replace(tzinfo=None)


def parse_export_filters(start_date, end_date, limit):
    start_day = _parse_export_date("开始日期", start_date)
    end_day = _parse_export_date("结束日期", end_date)
    if start_day is not None and end_day is not None and start_day > end_day:
        raise ExportRequestError("开始日期不能晚于结束日期。")

    if limit is None or (isinstance(limit, str) and not limit.strip()):
        parsed_limit = None
    elif isinstance(limit, bool):
        raise ExportRequestError("条数必须是 1 至 10,000 的整数。")
    elif isinstance(limit, int):
        parsed_limit = limit
    elif isinstance(limit, str):
        raw_limit = limit.strip()
        if not raw_limit.isascii() or not raw_limit.isdigit():
            raise ExportRequestError("条数必须是 1 至 10,000 的整数。")
        parsed_limit = int(raw_limit)
    else:
        raise ExportRequestError("条数必须是 1 至 10,000 的整数。")

    if parsed_limit is not None and not 1 <= parsed_limit <= MAX_EXPORT_ROWS:
        raise ExportRequestError("条数必须是 1 至 10,000 的整数。")

    return ExportFilters(
        start_at=(
            _local_midnight_to_utc_naive(start_day)
            if start_day is not None
            else None
        ),
        end_before=(
            _local_midnight_to_utc_naive(end_day + timedelta(days=1))
            if end_day is not None
            else None
        ),
        limit=parsed_limit,
    )


def _timestamp(value):
    return value.strftime("%Y-%m-%d %H:%M") if value is not None else ""


def _purchase_item_summary(order):
    return "；".join(
        f"{item.product.name} × {item.quantity}（¥{item.unit_price:.2f}）"
        for item in order.items
    )


def _sales_item_summary(order):
    return "；".join(
        f"{item.product.name} × {item.quantity}（¥{item.unit_price:.2f}）"
        for item in order.items
    )


def _apply_limit(query, filters):
    return query.limit(filters.limit) if filters.limit is not None else query


def _created_at_query(query, model, filters):
    if filters.start_at is not None:
        query = query.filter(model.created_at >= filters.start_at)
    if filters.end_before is not None:
        query = query.filter(model.created_at < filters.end_before)
    return query.order_by(model.created_at.desc(), model.id.desc())


def _product_rows(filters):
    return [
        (product.name, product.sku, product.purchase_price, product.sale_price, product.stock)
        for product in _apply_limit(
            Product.query.order_by(Product.id.desc()), filters
        ).all()
    ]


def _customer_rows(filters):
    return [
        (customer.name, customer.phone)
        for customer in _apply_limit(
            Customer.query.order_by(Customer.id.desc()), filters
        ).all()
    ]


def _supplier_rows(filters):
    return [
        (supplier.name, supplier.phone)
        for supplier in _apply_limit(
            Supplier.query.order_by(Supplier.id.desc()), filters
        ).all()
    ]


def _inventory_rows(filters):
    return [
        (product.name, product.sku, product.stock)
        for product in _apply_limit(
            Product.query.order_by(Product.name, Product.id), filters
        ).all()
    ]


def _inventory_transaction_rows(filters):
    query = InventoryTransaction.query.options(
        joinedload(InventoryTransaction.product)
    )
    transactions = _apply_limit(
        _created_at_query(query, InventoryTransaction, filters), filters
    ).all()
    return [
        (
            transaction.product.name,
            transaction.product.sku,
            MOVEMENT_LABELS.get(transaction.type, transaction.type),
            transaction.quantity,
            transaction.related_order_no,
            transaction.balance_after,
            _timestamp(transaction.created_at),
        )
        for transaction in transactions
    ]


def _purchase_order_rows(filters):
    query = PurchaseOrder.query.options(
        joinedload(PurchaseOrder.supplier),
        selectinload(PurchaseOrder.items).joinedload(PurchaseOrderItem.product),
    )
    orders = _apply_limit(
        _created_at_query(query, PurchaseOrder, filters), filters
    ).all()
    return [
        (
            order.order_no,
            order.supplier.name,
            STATUS_LABELS.get(order.status, order.status),
            order.total_amount,
            _timestamp(order.created_at),
            _purchase_item_summary(order),
        )
        for order in orders
    ]


def _sales_order_rows(filters):
    query = SalesOrder.query.options(
        joinedload(SalesOrder.customer),
        selectinload(SalesOrder.items).joinedload(SalesOrderItem.product),
    )
    orders = _apply_limit(
        _created_at_query(query, SalesOrder, filters), filters
    ).all()
    return [
        (
            order.order_no,
            order.customer.name,
            STATUS_LABELS.get(order.status, order.status),
            order.total_amount,
            _timestamp(order.created_at),
            _sales_item_summary(order),
        )
        for order in orders
    ]


def _receivable_rows(filters):
    query = AccountReceivable.query.options(
        joinedload(AccountReceivable.customer),
        joinedload(AccountReceivable.sales_order),
    )
    records = _apply_limit(
        _created_at_query(query, AccountReceivable, filters), filters
    ).all()
    return [
        (
            record.customer.name,
            record.sales_order.order_no,
            record.amount,
            STATUS_LABELS.get(record.status, record.status),
            _timestamp(record.created_at),
            _timestamp(record.paid_at),
        )
        for record in records
    ]


def _payable_rows(filters):
    query = AccountPayable.query.options(
        joinedload(AccountPayable.supplier),
        joinedload(AccountPayable.purchase_order),
    )
    records = _apply_limit(
        _created_at_query(query, AccountPayable, filters), filters
    ).all()
    return [
        (
            record.supplier.name,
            record.purchase_order.order_no,
            record.amount,
            STATUS_LABELS.get(record.status, record.status),
            _timestamp(record.created_at),
            _timestamp(record.paid_at),
        )
        for record in records
    ]


DATASETS = {
    "products": ExportDataset(
        "products",
        "商品列表",
        (
            ExportColumn("商品名称"),
            ExportColumn("SKU"),
            ExportColumn("采购价", "money"),
            ExportColumn("销售价", "money"),
            ExportColumn("库存数量", "integer"),
        ),
        _product_rows,
    ),
    "customers": ExportDataset(
        "customers",
        "客户列表",
        (ExportColumn("客户名称"), ExportColumn("联系电话")),
        _customer_rows,
    ),
    "suppliers": ExportDataset(
        "suppliers",
        "供应商列表",
        (ExportColumn("供应商名称"), ExportColumn("联系电话")),
        _supplier_rows,
    ),
    "inventory": ExportDataset(
        "inventory",
        "当前库存",
        (
            ExportColumn("商品名称"),
            ExportColumn("SKU"),
            ExportColumn("库存数量", "integer"),
        ),
        _inventory_rows,
    ),
    "inventory_transactions": ExportDataset(
        "inventory_transactions",
        "库存流水",
        (
            ExportColumn("商品名称"),
            ExportColumn("SKU"),
            ExportColumn("类型"),
            ExportColumn("数量", "integer"),
            ExportColumn("关联订单"),
            ExportColumn("结存", "integer"),
            ExportColumn("发生时间"),
        ),
        _inventory_transaction_rows,
        True,
    ),
    "purchase_orders": ExportDataset(
        "purchase_orders",
        "采购订单",
        (
            ExportColumn("采购订单号"),
            ExportColumn("供应商"),
            ExportColumn("状态"),
            ExportColumn("总金额", "money"),
            ExportColumn("创建时间"),
            ExportColumn("商品明细"),
        ),
        _purchase_order_rows,
        True,
    ),
    "sales_orders": ExportDataset(
        "sales_orders",
        "销售订单",
        (
            ExportColumn("销售订单号"),
            ExportColumn("客户"),
            ExportColumn("状态"),
            ExportColumn("总金额", "money"),
            ExportColumn("创建时间"),
            ExportColumn("商品明细"),
        ),
        _sales_order_rows,
        True,
    ),
    "receivables": ExportDataset(
        "receivables",
        "应收账款",
        (
            ExportColumn("客户"),
            ExportColumn("销售订单号"),
            ExportColumn("金额", "money"),
            ExportColumn("状态"),
            ExportColumn("创建时间"),
            ExportColumn("收款时间"),
        ),
        _receivable_rows,
        True,
    ),
    "payables": ExportDataset(
        "payables",
        "应付账款",
        (
            ExportColumn("供应商"),
            ExportColumn("采购订单号"),
            ExportColumn("金额", "money"),
            ExportColumn("状态"),
            ExportColumn("创建时间"),
            ExportColumn("付款时间"),
        ),
        _payable_rows,
        True,
    ),
}


def list_export_datasets():
    return tuple(DATASETS)


def export_dataset_options():
    return tuple(
        (specification.dataset_id, specification.title, specification.supports_date_filter)
        for specification in DATASETS.values()
    )


def get_export_rows(dataset, filters=None):
    specification = DATASETS.get(dataset)
    if specification is None:
        raise ExportRequestError("不支持导出该数据集。")
    return specification, specification.query_rows(filters or ExportFilters())


def _literal_text_cell(worksheet, row, column, value):
    cell = worksheet.cell(row=row, column=column, value="" if value is None else str(value))
    cell.data_type = "s"
    return cell


def _write_xlsx_sheet(worksheet, specification, rows):
    worksheet.title = specification.title[:31]
    column_count = len(specification.columns)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=column_count)
    title_cell = worksheet.cell(row=1, column=1, value=specification.title)
    title_cell.font = Font(name="等线", size=15, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill("solid", fgColor="243B53")
    title_cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 28

    for index, column in enumerate(specification.columns, start=1):
        cell = worksheet.cell(row=2, column=index, value=column.label)
        cell.font = Font(name="等线", bold=True, color="243B53")
        cell.fill = PatternFill("solid", fgColor="E8EEF5")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = min(
            max(len(column.label) * 2 + 4, 14), 42
        )

    if not rows:
        worksheet.merge_cells(
            start_row=3,
            start_column=1,
            end_row=3,
            end_column=column_count,
        )
        _literal_text_cell(worksheet, 3, 1, "没有符合条件的记录")

    for row_index, row_values in enumerate(rows, start=3):
        for column_index, (column, value) in enumerate(
            zip(specification.columns, row_values), start=1
        ):
            if column.kind == "text":
                cell = _literal_text_cell(worksheet, row_index, column_index, value)
            else:
                cell = worksheet.cell(row=row_index, column=column_index, value=value)
                if column.kind == "money":
                    cell.number_format = MONEY_FORMAT
                    cell.alignment = Alignment(horizontal="right")
            if column.kind == "text":
                worksheet.column_dimensions[cell.column_letter].width = min(
                    max(
                        worksheet.column_dimensions[cell.column_letter].width,
                        len(str(value or "")) + 3,
                    ),
                    42,
                )
            if column.kind == "text" and column.label == "商品明细":
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    worksheet.freeze_panes = "A3"
    last_row = max(2, worksheet.max_row)
    last_column = worksheet.cell(row=last_row, column=column_count).coordinate
    worksheet.auto_filter.ref = f"A2:{last_column}"
    return worksheet


def _render_xlsx(specification, rows):
    workbook = Workbook()
    _write_xlsx_sheet(workbook.active, specification, rows)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _render_all_xlsx(results):
    workbook = Workbook()
    for index, (specification, rows) in enumerate(results):
        worksheet = workbook.active if index == 0 else workbook.create_sheet()
        _write_xlsx_sheet(worksheet, specification, rows)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _table_for_dataset(specification, rows, available_width, cell_style, header_style):
    table_rows = [
        [Paragraph(escape(column.label), header_style) for column in specification.columns]
    ]
    table_rows.extend(
        [
            Paragraph(escape("" if value is None else str(value)), cell_style)
            for value in row_values
        ]
        for row_values in rows
    )
    table = Table(
        table_rows,
        colWidths=[available_width / len(specification.columns)] * len(specification.columns),
        repeatRows=1,
        hAlign="LEFT",
    )
    table_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#243B53")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    if rows:
        table_styles.append(
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")])
        )
    table.setStyle(TableStyle(table_styles))
    return table


def _render_pdf_sections(results, report_title=None):
    font_name = _register_pdf_font()
    page_size = landscape(A4) if any(
        len(specification.columns) > 5 for specification, _ in results
    ) else A4
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=page_size,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=report_title or results[0][0].title,
    )
    title_style = ParagraphStyle(
        "ExportTitle",
        fontName=font_name,
        fontSize=15,
        leading=20,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#243B53"),
        spaceAfter=8 * mm,
    )
    section_title_style = ParagraphStyle(
        "ExportSectionTitle",
        parent=title_style,
        fontSize=12,
        leading=16,
        alignment=TA_LEFT,
        spaceAfter=4 * mm,
    )
    cell_style = ParagraphStyle(
        "ExportCell",
        fontName=font_name,
        fontSize=7.5,
        leading=10,
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    header_style = ParagraphStyle(
        "ExportHeader",
        parent=cell_style,
        fontSize=8,
        leading=11,
        alignment=TA_CENTER,
        textColor=colors.white,
    )
    note_style = ParagraphStyle(
        "ExportSnapshotNote",
        parent=cell_style,
        fontSize=8,
        textColor=colors.HexColor("#52606D"),
        spaceAfter=2 * mm,
    )
    available_width = page_size[0] - document.leftMargin - document.rightMargin
    story = []
    if report_title:
        story.append(Paragraph(escape(report_title), title_style))
    for index, (specification, rows) in enumerate(results):
        if index:
            story.append(PageBreak())
        story.append(Paragraph(escape(specification.title), section_title_style))
        if not specification.supports_date_filter:
            story.append(Paragraph("当前快照，日期范围不适用", note_style))
        if rows:
            story.append(
                _table_for_dataset(
                    specification, rows, available_width, cell_style, header_style
                )
            )
        else:
            story.append(Paragraph("没有符合条件的记录", cell_style))
    document.build(story)
    return output.getvalue()


def _render_pdf(specification, rows):
    return _render_pdf_sections([(specification, rows)])


def _render_all_pdf(results):
    return _render_pdf_sections(results, "ERP 全量数据导出")


def generate_export(
    dataset,
    file_format,
    *,
    start_date=None,
    end_date=None,
    limit=None,
):
    if file_format not in {"xlsx", "pdf"}:
        raise ExportRequestError("导出格式仅支持 Excel 或 PDF。")
    filters = parse_export_filters(start_date, end_date, limit)
    if dataset == ALL_DATASET_ID:
        results = [get_export_rows(key, filters) for key in DATASETS]
    else:
        results = [get_export_rows(dataset, filters)]

    if file_format == "xlsx":
        content = (
            _render_all_xlsx(results)
            if dataset == ALL_DATASET_ID
            else _render_xlsx(*results[0])
        )
        extension = "xlsx"
        mime_type = XLSX_MIME_TYPE
    else:
        content = (
            _render_all_pdf(results)
            if dataset == ALL_DATASET_ID
            else _render_pdf(*results[0])
        )
        extension = "pdf"
        mime_type = PDF_MIME_TYPE

    today = datetime.now().strftime("%Y%m%d")
    filename_prefix = "all_data" if dataset == ALL_DATASET_ID else dataset
    return ExportFile(
        content=content,
        filename=f"{filename_prefix}_{today}.{extension}",
        mime_type=mime_type,
    )


def _register_pdf_font():
    font_name = "STSong-Light"
    try:
        pdfmetrics.getFont(font_name)
    except KeyError:
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    return font_name
