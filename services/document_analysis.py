from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from pypdf import PdfReader


MAX_DOCUMENT_BYTES = 10 * 1024 * 1024
MAX_DOCUMENT_PAGES = 50
MAX_WORKBOOK_CELLS = 20_000
MAX_WORKBOOK_SCANNED_CELLS = 100_000
MAX_EXTRACTED_CHARS = 50_000
MAX_XLSX_MEMBERS = 100
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024


class DocumentInputError(ValueError):
    """A safe, user-facing document validation or extraction error."""


def _validate_stream_size(stream):
    current_position = None
    try:
        current_position = stream.tell()
        stream.seek(0, 2)
        size = stream.tell()
    except (AttributeError, OSError, ValueError):
        raise DocumentInputError("无法读取上传文件。") from None
    finally:
        if current_position is not None:
            try:
                stream.seek(current_position)
            except (AttributeError, OSError, ValueError):
                pass
    if size > MAX_DOCUMENT_BYTES:
        raise DocumentInputError("文件过大，请上传 10 MiB 以内的文件。")


def _read_prefix(stream, length):
    try:
        position = stream.tell()
        stream.seek(0)
        prefix = stream.read(length)
        stream.seek(position)
    except (AttributeError, OSError, ValueError) as exc:
        raise DocumentInputError("无法读取上传文件。") from exc
    if isinstance(prefix, str):
        prefix = prefix.encode("utf-8", errors="replace")
    return prefix


def _append_limited(parts, value, current_length):
    if current_length + len(value) > MAX_EXTRACTED_CHARS:
        raise DocumentInputError("文件提取内容超过 50,000 个字符，请拆分文件后再上传。")
    parts.append(value)
    return current_length + len(value)


def _extract_pdf(stream):
    if not _read_prefix(stream, 5).startswith(b"%PDF-"):
        raise DocumentInputError("文件内容不是有效的 PDF。")

    try:
        stream.seek(0)
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise DocumentInputError("暂不支持加密 PDF，请先移除密码后再上传。")
        if len(reader.pages) > MAX_DOCUMENT_PAGES:
            raise DocumentInputError("PDF 页数超过 50 页，请拆分后再上传。")

        parts = []
        total_length = 0
        found_text = False
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            if page_text.strip():
                found_text = True
            section = f"\n[PDF 第 {page_number} 页]\n{page_text}\n"
            total_length = _append_limited(parts, section, total_length)
        if not found_text:
            raise DocumentInputError(
                "PDF 中没有提取到文字；扫描版 PDF 暂不支持 OCR，请上传可搜索文本 PDF。"
            )
        return "".join(parts).strip()
    except DocumentInputError:
        raise
    except Exception as exc:
        raise DocumentInputError("PDF 文件损坏或格式无效。") from exc


def _cell_text(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        try:
            return value.isoformat(sep=" ")
        except TypeError:
            return value.isoformat()
    return str(value)


def _validate_xlsx_archive(stream):
    if not _read_prefix(stream, 4).startswith(b"PK"):
        raise DocumentInputError("文件内容不是有效的 XLSX 工作簿。")
    try:
        stream.seek(0)
        with ZipFile(stream) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_MEMBERS:
                raise DocumentInputError("XLSX 文件包含过多内部条目，无法安全处理。")
            expanded_size = sum(member.file_size for member in members)
            if expanded_size > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise DocumentInputError("XLSX 解压后内容过大，请缩小文件后再上传。")
            member_names = {member.filename for member in members}
            if "[Content_Types].xml" not in member_names or "xl/workbook.xml" not in member_names:
                raise DocumentInputError("文件内容不是有效的 XLSX 工作簿。")
    except DocumentInputError:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise DocumentInputError("XLSX 文件损坏或格式无效。") from exc


def _extract_xlsx(stream):
    _validate_xlsx_archive(stream)
    workbook = None
    try:
        stream.seek(0)
        workbook = load_workbook(
            stream,
            read_only=True,
            data_only=True,
            keep_links=False,
        )
        parts = []
        total_length = 0
        cell_count = 0
        scanned_cell_count = 0
        for worksheet in workbook.worksheets:
            sheet_has_content = False
            worksheet.reset_dimensions()
            for row in worksheet.iter_rows(values_only=True):
                scanned_cell_count += len(row)
                if scanned_cell_count > MAX_WORKBOOK_SCANNED_CELLS:
                    raise DocumentInputError(
                        "工作簿实际表格范围超过 100,000 个单元格，请精简后再上传。"
                    )
                non_empty_values = []
                for value in row:
                    text = _cell_text(value)
                    if text:
                        cell_count += 1
                        if cell_count > MAX_WORKBOOK_CELLS:
                            raise DocumentInputError(
                                "工作簿非空单元格超过 20,000 个，请拆分后再上传。"
                            )
                    non_empty_values.append(text)
                if any(non_empty_values):
                    if not sheet_has_content:
                        total_length = _append_limited(
                            parts, f"\n[工作表：{worksheet.title}]\n", total_length
                        )
                        sheet_has_content = True
                    line = "\t".join(non_empty_values) + "\n"
                    total_length = _append_limited(parts, line, total_length)

        extracted = "".join(parts).strip()
        if not extracted:
            raise DocumentInputError("XLSX 工作簿中没有可提取的文字或数值。")
        return extracted
    except DocumentInputError:
        raise
    except Exception as exc:
        raise DocumentInputError("XLSX 文件损坏或格式无效。") from exc
    finally:
        if workbook is not None:
            workbook.close()


def extract_document_text(filename, stream):
    if not isinstance(filename, str) or not filename.strip():
        raise DocumentInputError("请先选择 PDF 或 XLSX 文件。")
    extension = Path(Path(filename).name).suffix.casefold()
    if extension not in {".pdf", ".xlsx"}:
        raise DocumentInputError("仅支持 PDF 和 .xlsx 文件；旧版 .xls 暂不支持。")

    _validate_stream_size(stream)
    if extension == ".pdf":
        return _extract_pdf(stream)
    return _extract_xlsx(stream)
